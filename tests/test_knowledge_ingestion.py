from pathlib import Path

import pytest
from core.knowledge.ingestion import (
    build_provenance_metadata,
    extract_text_from_docx,
    extract_text_from_text_file,
    ingest_document,
    ingest_into_vector_store,
)

from core.knowledge.vector_store import VectorStore

from docx import Document


def test_ingest_text_file(tmp_path: Path):
    document = tmp_path / "notes.txt"
    document.write_text(
        "Kraken is my Neptune 4 printer.",
        encoding="utf-8",
    )

    result = ingest_document(document)

    assert result["source_type"] == "text"
    assert result["filename"] == "notes.txt"
    assert result["full_text"] == "Kraken is my Neptune 4 printer."
    assert result["text_length"] == len(result["full_text"])

def test_reingestion_with_empty_document_removes_old_chunks(
    tmp_path: Path,
    monkeypatch,
):
    document = tmp_path / "empty-update.txt"
    store = VectorStore(tmp_path / "vectors")

    monkeypatch.setattr(
        "core.knowledge.ingestion.create_embeddings",
        lambda texts: [[0.1, 0.2, 0.3] for _ in texts],
    )

    document.write_text(
        "Snowball has knowledge.",
        encoding="utf-8",
    )

    ingest_into_vector_store(
        document,
        store,
        document_id="empty-update",
    )

    assert store.count() == 1

    document.write_text("", encoding="utf-8")

    result = ingest_into_vector_store(
        document,
        store,
        document_id="empty-update",
    )

    assert result["chunk_count"] == 0
    assert result["vector_ids"] == []
    assert store.count() == 0

    
def test_ingest_markdown_file(tmp_path: Path):
    document = tmp_path / "snowball.md"
    document.write_text(
        "# Snowball\n\nPersistent personal AI.",
        encoding="utf-8",
    )

    result = ingest_document(document)

    assert result["source_type"] == "text"
    assert result["filename"] == "snowball.md"
    assert "Persistent personal AI." in result["full_text"]


def test_extract_text_file_requires_existing_file(tmp_path: Path):
    missing = tmp_path / "missing.txt"

    with pytest.raises(FileNotFoundError):
        extract_text_from_text_file(missing)


def test_ingest_rejects_unsupported_document_type(tmp_path: Path):
    document = tmp_path / "binary.xyz"
    document.write_text("test", encoding="utf-8")

    with pytest.raises(ValueError):
        ingest_document(document)


def test_ingest_rejects_file_without_extension(tmp_path: Path):
    document = tmp_path / "README"
    document.write_text("Snowball", encoding="utf-8")

    with pytest.raises(ValueError):
        ingest_document(document)


def test_ingest_into_vector_store(tmp_path: Path):
    from core.knowledge.vector_store import VectorStore

    document = tmp_path / "printers.txt"

    document.write_text(
        (
            "Kraken is my Neptune 4 3D printer. "
            "Ferris is my Ender 3 S1 Pro 3D printer."
        ),
        encoding="utf-8",
    )

    store = VectorStore(tmp_path / "vectors")

    result = ingest_into_vector_store(
        document,
        store,
        document_id="printer-notes",
        chunk_size=50,
        overlap=10,
    )

    assert result["document_id"] == "printer-notes"
    assert result["filename"] == "printers.txt"
    assert result["source_type"] == "text"
    assert result["chunk_count"] > 0
    assert len(result["vector_ids"]) == result["chunk_count"]

    assert store.count() == result["chunk_count"]

def test_reingestion_replaces_stale_document_chunks(
    tmp_path: Path,
    monkeypatch,
):
    document = tmp_path / "changing-notes.txt"
    store = VectorStore(tmp_path / "vectors")

    # Keep the test deterministic and independent of Ollama.
    monkeypatch.setattr(
        "core.knowledge.ingestion.create_embeddings",
        lambda texts: [[0.1, 0.2, 0.3] for _ in texts],
    )

    # Version 1 is large enough to create multiple chunks.
    document.write_text(
        "A" * 250,
        encoding="utf-8",
    )

    first_result = ingest_into_vector_store(
        document,
        store,
        document_id="changing-document",
        chunk_size=100,
        overlap=0,
    )

    assert first_result["chunk_count"] == 3
    assert store.count() == 3

    # Version 2 is much shorter and should replace version 1.
    document.write_text(
        "Snowball updated knowledge.",
        encoding="utf-8",
    )

    second_result = ingest_into_vector_store(
        document,
        store,
        document_id="changing-document",
        chunk_size=100,
        overlap=0,
    )

    assert second_result["chunk_count"] == 1
    assert store.count() == 1

    stored = store.collection.get(
        where={"document_id": "changing-document"}
    )

    assert stored["ids"] == ["changing-document_chunk_0"]
    assert stored["documents"] == ["Snowball updated knowledge."]


def test_ingested_document_can_be_retrieved_semantically(
    tmp_path: Path,
):
    from core.knowledge.vector_store import VectorStore
    from core.memory.retrieval import retrieve_relevant_chunks

    document = tmp_path / "workshop-notes.txt"

    document.write_text(
        (
            "Kraken is my Neptune 4 3D printer. "
            "I use Kraken for making physical printed parts.\n\n"
            "Monty is my snake plant. "
            "Sensors will eventually let Monty react to his environment.\n\n"
            "Brain is PeenQi's clockwork mouse companion."
        ),
        encoding="utf-8",
    )

    store = VectorStore(tmp_path / "vectors")

    ingest_into_vector_store(
        document,
        store,
        document_id="workshop-notes",
        chunk_size=100,
        overlap=20,
    )

    results = retrieve_relevant_chunks(
        "Which machine can make physical objects for me?",
        vector_store=store,
        result_count=1,
    )

    assert len(results) == 1
    assert "Kraken" in results[0]["text"]


def test_build_provenance_metadata():
    metadata = build_provenance_metadata(
        source_type="snowball_project_document",
        authority="reference",
        project="Snowball AI/OS",
        source_date="2025-09-30",
        original_source="historical Snowball documentation",
    )

    assert metadata == {
        "provenance_source_type": "snowball_project_document",
        "authority": "reference",
        "project": "Snowball AI/OS",
        "source_date": "2025-09-30",
        "original_source": "historical Snowball documentation",
    }


def test_build_provenance_metadata_omits_optional_empty_fields():
    metadata = build_provenance_metadata(
        source_type="snowball_project_document",
        authority="reference",
    )

    assert metadata == {
        "provenance_source_type": "snowball_project_document",
        "authority": "reference",
    }


def test_ingestion_preserves_provenance_metadata(
    tmp_path: Path,
    monkeypatch,
):
    document_path = tmp_path / "snowball-history.txt"
    document_path.write_text(
        "Snowball began as a personal AI project.",
        encoding="utf-8",
    )

    vector_store = VectorStore(
        tmp_path / "vectors"
    )

    provenance = build_provenance_metadata(
        source_type="snowball_project_document",
        authority="reference",
        project="Snowball AI/OS",
        source_date="2025-09-30",
        original_source="historical Snowball documentation",
    )

    # Keep this test deterministic and independent of Ollama.
    monkeypatch.setattr(
        "core.knowledge.ingestion.create_embeddings",
        lambda texts: [[0.1, 0.2, 0.3] for _ in texts],
    )

    result = ingest_into_vector_store(
        document_path,
        vector_store,
        metadata=provenance,
    )

    assert result["chunk_count"] == 1

    matches = vector_store.search(
        [0.1, 0.2, 0.3],
        result_count=1,
    )

    assert len(matches) == 1

    metadata = matches[0]["metadata"]

    assert metadata["source_type"] == "text"
    assert (
        metadata["provenance_source_type"]
        == "snowball_project_document"
    )
    assert metadata["authority"] == "reference"
    assert metadata["project"] == "Snowball AI/OS"
    assert metadata["source_date"] == "2025-09-30"
    assert (
        metadata["original_source"]
        == "historical Snowball documentation"
    )
    assert metadata["filename"] == "snowball-history.txt"

def test_extract_text_from_docx(tmp_path: Path):
    path = tmp_path / "example.docx"

    document = Document()
    document.add_paragraph("Snowball AI/OS")
    document.add_paragraph("Unified project knowledge")
    document.save(path)

    result = extract_text_from_docx(path)

    assert result["source_type"] == "docx"
    assert result["filename"] == "example.docx"
    assert "Snowball AI/OS" in result["full_text"]
    assert "Unified project knowledge" in result["full_text"]


def test_extract_text_from_docx_requires_existing_file(tmp_path: Path):
    missing = tmp_path / "missing.docx"

    with pytest.raises(FileNotFoundError):
        extract_text_from_docx(missing)


def test_extract_text_from_docx_preserves_paragraph_and_table_order(tmp_path: Path):
    path = tmp_path / "ordered.docx"

    document = Document()
    document.add_paragraph("Before table")

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Snowball"
    table.cell(1, 1).text = "AI/OS"

    document.add_paragraph("After table")
    document.save(path)

    result = extract_text_from_docx(path)
    text = result["full_text"]

    assert "Name | Value" in text
    assert "Snowball | AI/OS" in text

    assert text.index("Before table") < text.index("Name | Value")
    assert text.index("Name | Value") < text.index("Snowball | AI/OS")
    assert text.index("Snowball | AI/OS") < text.index("After table")